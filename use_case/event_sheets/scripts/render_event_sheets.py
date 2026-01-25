#!/usr/bin/env python3
"""
이벤트 시트 렌더링 스크립트

story-template 시트를 event_name 기준으로 복사하고 렌더링합니다.

동작:
1. template.xlsx를 output.xlsx로 복사
2. data.json에서 events를 event_name 기준으로 그룹화
3. 각 그룹별로:
   - story-template 시트를 복사하여 새 시트 생성
   - 시트명: event_title들 조합 (예: "1부 대면 팬사인회 2부 포토회")
   - 해당 시트에 그룹의 이벤트 데이터로 렌더링
4. story-template 시트 삭제
5. 저장
"""

import sys
import json
import shutil
from pathlib import Path
from collections import OrderedDict
from typing import Dict, List, Any

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl이 필요합니다: pip install openpyxl")
    sys.exit(1)

# xlsx_template_renderer import를 위해 프로젝트 루트 추가
SCRIPT_DIR = Path(__file__).parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))

try:
    from xlsx_template_renderer import render_template
except ImportError:
    print("xlsx_template_renderer를 찾을 수 없습니다.")
    print(f"PROJECT_ROOT: {PROJECT_ROOT}")
    sys.exit(1)


TEMPLATE_SHEET_NAME = "story-template"


def load_data(data_path: Path) -> Dict[str, Any]:
    """data.json 로드"""
    with open(data_path, "r", encoding="utf-8") as f:
        return json.load(f)


def group_events_by_name(events: List[Dict]) -> Dict[str, List[Dict]]:
    """
    events를 event_name 기준으로 그룹화
    
    Returns:
        OrderedDict: {event_name: [events...]} (순서 유지)
    """
    groups = OrderedDict()
    
    for event in events:
        event_name = event.get("overview", {}).get("event_name", "Unknown")
        if event_name not in groups:
            groups[event_name] = []
        groups[event_name].append(event)
    
    return groups


def sanitize_sheet_name(name: str) -> str:
    """
    시트명에서 Excel이 허용하지 않는 문자 제거
    
    Excel 시트명 금지 문자: : \ / ? * [ ]
    """
    invalid_chars = [':', '\\', '/', '?', '*', '[', ']']
    for char in invalid_chars:
        name = name.replace(char, '')
    return name


def get_sheet_name(events_group: List[Dict]) -> str:
    """
    이벤트 그룹에서 시트명 생성
    
    단일 이벤트: event_title 그대로
    복합 이벤트: event_title들 공백으로 조합
    """
    titles = [e.get("event_title", "Unknown") for e in events_group]
    name = " ".join(titles)
    return sanitize_sheet_name(name)


def copy_sheet(wb, source_name: str, target_name: str):
    """
    시트 복사
    
    Args:
        wb: Workbook
        source_name: 원본 시트명
        target_name: 새 시트명
    """
    if source_name not in wb.sheetnames:
        raise ValueError(f"원본 시트 '{source_name}'를 찾을 수 없습니다")
    
    source = wb[source_name]
    target = wb.copy_worksheet(source)
    target.title = target_name
    
    return target


def render_event_sheets(
    template_path: Path,
    data_path: Path,
    output_path: Path,
    verbose: bool = False
) -> List[str]:
    """
    이벤트 시트 렌더링
    
    Args:
        template_path: 템플릿 xlsx 파일 경로
        data_path: data.json 파일 경로
        output_path: 출력 xlsx 파일 경로
        verbose: 상세 출력 여부
    
    Returns:
        생성된 시트명 목록
    """
    # 1. 템플릿 복사
    shutil.copy2(template_path, output_path)
    
    # 2. data.json 로드
    data = load_data(data_path)
    events = data.get("events", [])
    
    if not events:
        if verbose:
            print("  경고: events가 비어있습니다")
        return []
    
    # 3. events를 event_name 기준으로 그룹화
    groups = group_events_by_name(events)
    
    if verbose:
        print(f"  이벤트 그룹: {len(groups)}개")
        for event_name, group_events in groups.items():
            print(f"    - {event_name}: {len(group_events)}개 이벤트")
    
    # 4. 워크북 열기 (시트 복사용)
    wb = load_workbook(output_path)
    
    if TEMPLATE_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"템플릿 시트 '{TEMPLATE_SHEET_NAME}'를 찾을 수 없습니다")
    
    # 5. 각 그룹별로 시트 복사 및 렌더링
    created_sheets = []
    
    for event_name, group_events in groups.items():
        # 시트명 생성
        sheet_name = get_sheet_name(group_events)
        
        # 시트명 길이 제한 (Excel은 31자 제한)
        if len(sheet_name) > 31:
            sheet_name = sheet_name[:28] + "..."
        
        if verbose:
            print(f"  시트 생성: '{sheet_name}'")
        
        # 시트 복사
        copy_sheet(wb, TEMPLATE_SHEET_NAME, sheet_name)
        created_sheets.append(sheet_name)
    
    # 6. story-template 시트 삭제
    if TEMPLATE_SHEET_NAME in wb.sheetnames:
        del wb[TEMPLATE_SHEET_NAME]
    
    # 7. 워크북 저장 (시트 복사 적용)
    wb.save(output_path)
    wb.close()
    
    # 8. 각 시트별로 렌더링
    for idx, (event_name, group_events) in enumerate(groups.items()):
        sheet_name = created_sheets[idx]
        
        # 해당 그룹의 events만 포함된 데이터 생성
        group_data = {
            **data,
            "events": group_events
        }
        
        if verbose:
            print(f"  렌더링: '{sheet_name}' ({len(group_events)}개 이벤트)")
        
        # 해당 시트만 렌더링
        render_template(str(output_path), group_data, sheets=[sheet_name])
    
    return created_sheets


def main():
    """CLI 진입점"""
    import argparse
    
    parser = argparse.ArgumentParser(
        description="이벤트 시트 렌더링 스크립트"
    )
    parser.add_argument(
        "template",
        help="템플릿 xlsx 파일 경로"
    )
    parser.add_argument(
        "data",
        help="data.json 파일 경로"
    )
    parser.add_argument(
        "-o", "--output",
        required=True,
        help="출력 파일 경로"
    )
    parser.add_argument(
        "-v", "--verbose",
        action="store_true",
        help="상세 출력"
    )
    
    args = parser.parse_args()
    
    template_path = Path(args.template)
    data_path = Path(args.data)
    output_path = Path(args.output)
    
    # 파일 존재 확인
    if not template_path.exists():
        print(f"오류: 템플릿 파일을 찾을 수 없습니다: {template_path}")
        sys.exit(1)
    
    if not data_path.exists():
        print(f"오류: 데이터 파일을 찾을 수 없습니다: {data_path}")
        sys.exit(1)
    
    try:
        if args.verbose:
            print(f"템플릿: {template_path}")
            print(f"데이터: {data_path}")
            print(f"출력: {output_path}")
        
        sheets = render_event_sheets(
            template_path,
            data_path,
            output_path,
            verbose=args.verbose
        )
        
        print(f"완료: {output_path}")
        if args.verbose:
            print(f"생성된 시트: {', '.join(sheets)}")
        
    except Exception as e:
        print(f"오류: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()

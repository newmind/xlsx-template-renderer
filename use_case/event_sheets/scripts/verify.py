#!/usr/bin/env python3
"""
output.xlsx 검증 스크립트

검증 항목:
1. 파일 존재 여부
2. 시트명 검증 (data.json 기반 예상 시트명과 비교)
3. 템플릿 변수 잔존 확인 ({{, {%)
4. 에러 메시지 검출
"""

import sys
import json
import re
from pathlib import Path
from collections import OrderedDict
from typing import List, Dict, Any

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl이 필요합니다: pip install openpyxl")
    sys.exit(1)


# 렌더링에서 제외되는 시트들 (템플릿에 원래 있던 시트)
EXCLUDED_SHEETS = [
    "Overview", "Rule book", "story-template", "안내메일",
    "▶️ 대면", "팬사인회", "포토회", "게임회", "베이커리 이벤트",
    "▶️ 스페셜 이벤트", "일일점장", "스페셜 이벤트(가챠)", "스페셜 이벤트(가챠 외)",
    "▶️ 영상통화 이벤트", "영상통화", "유닛 영상통화", "개인 영상통화", "11 온라인 투샷회"
]


class Colors:
    RED = '\033[0;31m'
    GREEN = '\033[0;32m'
    YELLOW = '\033[1;33m'
    CYAN = '\033[0;36m'
    NC = '\033[0m'


def load_data(data_path: Path) -> Dict[str, Any]:
    """data.json 로드"""
    with open(data_path, "r", encoding="utf-8") as f:
        return json.load(f)


def sanitize_sheet_name(name: str) -> str:
    """
    시트명에서 Excel이 허용하지 않는 문자 제거
    
    Excel 시트명 금지 문자: : \\ / ? * [ ]
    """
    invalid_chars = [':', '\\', '/', '?', '*', '[', ']']
    for char in invalid_chars:
        name = name.replace(char, '')
    return name


def get_expected_sheets(data: Dict[str, Any]) -> List[str]:
    """
    data.json에서 예상 시트명 계산
    
    render_event_sheets.py와 동일한 로직 사용
    """
    events = data.get("events", [])
    if not events:
        return []
    
    # event_name 기준 그룹화
    groups = OrderedDict()
    for event in events:
        event_name = event.get("overview", {}).get("event_name", "Unknown")
        if event_name not in groups:
            groups[event_name] = []
        groups[event_name].append(event)
    
    # 각 그룹별 시트명 생성
    sheet_names = []
    for event_name, group_events in groups.items():
        titles = [e.get("event_title", "Unknown") for e in group_events]
        sheet_name = " ".join(titles)
        sheet_name = sanitize_sheet_name(sheet_name)
        # 시트명 길이 제한 (Excel은 31자 제한)
        if len(sheet_name) > 31:
            sheet_name = sheet_name[:28] + "..."
        sheet_names.append(sheet_name)
    
    return sheet_names


def verify_case(case_dir: Path, case_name: str) -> dict:
    """단일 케이스 검증"""
    result = {
        "case": case_name,
        "status": "PASS",
        "errors": [],
        "warnings": [],
        "sheets": [],
        "expected_sheets": [],
    }
    
    output_file = case_dir / "output.xlsx"
    data_file = case_dir / "data.json"
    
    # 1. 파일 존재 확인
    if not output_file.exists():
        result["status"] = "FAIL"
        result["errors"].append("output.xlsx 파일이 없습니다")
        return result
    
    if not data_file.exists():
        result["status"] = "FAIL"
        result["errors"].append("data.json 파일이 없습니다")
        return result
    
    # 2. 예상 시트명 계산
    data = load_data(data_file)
    expected_sheets = get_expected_sheets(data)
    result["expected_sheets"] = expected_sheets
    
    # 3. 워크북 로드
    try:
        wb = load_workbook(output_file, data_only=True)
    except Exception as e:
        result["status"] = "FAIL"
        result["errors"].append(f"파일 로드 실패: {e}")
        return result
    
    # 4. 시트명 확인
    actual_sheets = wb.sheetnames
    result["sheets"] = actual_sheets
    
    # 렌더링된 시트만 추출 (템플릿 원본 시트 제외)
    rendered_sheets = [s for s in actual_sheets if s not in EXCLUDED_SHEETS]
    
    # 예상 시트 존재 확인
    for exp_sheet in expected_sheets:
        if exp_sheet not in rendered_sheets:
            result["status"] = "FAIL"
            result["errors"].append(f"예상 시트 없음: '{exp_sheet}'")
    
    # 예상하지 않은 시트 확인
    for actual_sheet in rendered_sheets:
        if actual_sheet not in expected_sheets:
            result["warnings"].append(f"예상하지 않은 시트: '{actual_sheet}'")
    
    # 5. 템플릿 변수 잔존 확인 (렌더링된 시트만)
    template_pattern = re.compile(r'\{\{|\{%')
    error_pattern = re.compile(r'#ERROR|#VALUE|#NAME|#REF')
    
    for sheet_name in rendered_sheets:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    # 템플릿 변수 체크
                    if template_pattern.search(cell.value):
                        result["warnings"].append(
                            f"[{sheet_name}] {cell.coordinate}: 템플릿 변수 잔존 - '{cell.value[:50]}...'"
                        )
                    # 에러 메시지 체크
                    if error_pattern.search(cell.value):
                        result["errors"].append(
                            f"[{sheet_name}] {cell.coordinate}: 에러 발견 - '{cell.value}'"
                        )
    
    wb.close()
    
    # 에러가 있으면 FAIL
    if result["errors"]:
        result["status"] = "FAIL"
    elif result["warnings"]:
        result["status"] = "WARN"
    
    return result


def print_result(result: dict):
    """결과 출력"""
    case = result["case"]
    status = result["status"]
    
    if status == "PASS":
        status_str = f"{Colors.GREEN}PASS{Colors.NC}"
    elif status == "WARN":
        status_str = f"{Colors.YELLOW}WARN{Colors.NC}"
    else:
        status_str = f"{Colors.RED}FAIL{Colors.NC}"
    
    # 렌더링된 시트만 표시
    rendered = [s for s in result["sheets"] if s not in EXCLUDED_SHEETS]
    sheets_info = ", ".join(rendered) if rendered else "없음"
    
    print(f"[{status_str}] {case}")
    print(f"       예상: {result['expected_sheets']}")
    print(f"       실제: [{sheets_info}]")
    
    for error in result["errors"]:
        print(f"       {Colors.RED}ERROR: {error}{Colors.NC}")
    
    for warning in result["warnings"][:5]:  # 경고는 최대 5개만 표시
        print(f"       {Colors.YELLOW}WARN: {warning}{Colors.NC}")
    
    if len(result["warnings"]) > 5:
        print(f"       {Colors.YELLOW}... 외 {len(result['warnings']) - 5}개 경고{Colors.NC}")


def main():
    script_dir = Path(__file__).parent
    base_dir = script_dir.parent
    cases_dir = base_dir / "cases"
    
    # 특정 케이스만 검증
    if len(sys.argv) > 1:
        case_names = sys.argv[1:]
    else:
        # 전체 케이스
        case_names = sorted([d.name for d in cases_dir.iterdir() if d.is_dir()])
    
    print("========================================")
    print("이벤트 시트 검증")
    print("========================================")
    print()
    
    pass_count = 0
    warn_count = 0
    fail_count = 0
    
    for case_name in case_names:
        case_dir = cases_dir / case_name
        if not case_dir.exists():
            print(f"{Colors.RED}[ERROR] 케이스 없음: {case_name}{Colors.NC}")
            fail_count += 1
            continue
        
        result = verify_case(case_dir, case_name)
        print_result(result)
        print()
        
        if result["status"] == "PASS":
            pass_count += 1
        elif result["status"] == "WARN":
            warn_count += 1
        else:
            fail_count += 1
    
    print("========================================")
    print(f"결과: {Colors.GREEN}PASS: {pass_count}{Colors.NC}, "
          f"{Colors.YELLOW}WARN: {warn_count}{Colors.NC}, "
          f"{Colors.RED}FAIL: {fail_count}{Colors.NC}")
    print("========================================")
    
    # FAIL이 있으면 exit code 1
    sys.exit(1 if fail_count > 0 else 0)


if __name__ == "__main__":
    main()

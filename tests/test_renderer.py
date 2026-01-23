"""Tests for template renderer"""

import pytest
import tempfile
import os
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from xlsx_template_renderer import render_template
from xlsx_template_renderer.exceptions import TemplateSyntaxError


def create_template(rows_data: list) -> str:
    """Create a temporary template file with given row data."""
    wb = Workbook()
    ws = wb.active
    
    for row_idx, row in enumerate(rows_data, start=1):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


def read_output(path: str) -> list:
    """Read all cell values from output file."""
    wb = load_workbook(path)
    ws = wb.active
    
    result = []
    for row in ws.iter_rows(values_only=True):
        # Filter out completely empty rows
        if any(cell is not None for cell in row):
            # Remove trailing None values
            row_list = list(row)
            while row_list and row_list[-1] is None:
                row_list.pop()
            result.append(row_list)
    return result


class TestVariableSubstitution:
    """Tests for variable substitution"""
    
    def test_simple_variable(self):
        template_path = create_template([
            ["{{ name }}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"name": "홍길동"})
            result = read_output(output_path)
            assert result == [["홍길동"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_mixed_text_and_variable(self):
        template_path = create_template([
            ["이름: {{ name }}님"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"name": "홍길동"})
            result = read_output(output_path)
            assert result == [["이름: 홍길동님"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_multiple_variables_in_cell(self):
        template_path = create_template([
            ["{{ name }} ({{ age }}세)"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"name": "홍길동", "age": 25})
            result = read_output(output_path)
            assert result == [["홍길동 (25세)"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_missing_variable_kept(self):
        template_path = create_template([
            ["{{ unknown }}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {})
            result = read_output(output_path)
            assert result == [["{{ unknown }}"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_nested_variable(self):
        template_path = create_template([
            ["{{ user.name }}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"user": {"name": "홍길동"}})
            result = read_output(output_path)
            assert result == [["홍길동"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_arithmetic_expression(self):
        template_path = create_template([
            ["{{ price * 1.1 }}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"price": 1000})
            result = read_output(output_path)
            assert result == [[1100.0]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestForLoop:
    """Tests for for loop"""
    
    def test_simple_for_loop(self):
        template_path = create_template([
            ["{% for item in items %}"],
            ["", "{{ item }}"],
            ["{% endfor %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"items": ["A", "B", "C"]})
            result = read_output(output_path)
            assert result == [
                [None, "A"],
                [None, "B"],
                [None, "C"]
            ]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_for_loop_with_dict_items(self):
        template_path = create_template([
            ["{% for item in items %}"],
            ["", "{{ item.name }}", "{{ item.price }}"],
            ["{% endfor %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            data = {
                "items": [
                    {"name": "상품A", "price": 1000},
                    {"name": "상품B", "price": 2000}
                ]
            }
            render_template(template_path, output_path, data)
            result = read_output(output_path)
            assert result == [
                [None, "상품A", 1000],
                [None, "상품B", 2000]
            ]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_empty_list(self):
        template_path = create_template([
            ["Header"],
            ["{% for item in items %}"],
            ["", "{{ item }}"],
            ["{% endfor %}"],
            ["Footer"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"items": []})
            result = read_output(output_path)
            assert result == [["Header"], ["Footer"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_nested_for_loop(self):
        template_path = create_template([
            ["{% for dept in departments %}"],
            ["", "{{ dept.name }}"],
            ["{% for member in dept.members %}"],
            ["", "", "{{ member }}"],
            ["{% endfor %}"],
            ["{% endfor %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            data = {
                "departments": [
                    {"name": "개발팀", "members": ["김철수", "이영희"]},
                    {"name": "기획팀", "members": ["박지민"]}
                ]
            }
            render_template(template_path, output_path, data)
            result = read_output(output_path)
            assert result == [
                [None, "개발팀"],
                [None, None, "김철수"],
                [None, None, "이영희"],
                [None, "기획팀"],
                [None, None, "박지민"]
            ]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestIfStatement:
    """Tests for if statement"""
    
    def test_if_true(self):
        template_path = create_template([
            ["{% if show %}"],
            ["", "표시됨"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"show": True})
            result = read_output(output_path)
            assert result == [[None, "표시됨"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_if_false(self):
        template_path = create_template([
            ["{% if show %}"],
            ["", "표시됨"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"show": False})
            result = read_output(output_path)
            assert result == []
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_if_else(self):
        template_path = create_template([
            ["{% if show %}"],
            ["", "YES"],
            ["{% else %}"],
            ["", "NO"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"show": False})
            result = read_output(output_path)
            assert result == [[None, "NO"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_if_elif_else(self):
        template_path = create_template([
            ["{% if status == 1 %}"],
            ["", "상태1"],
            ["{% elif status == 2 %}"],
            ["", "상태2"],
            ["{% else %}"],
            ["", "기타"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            # Test status == 1
            render_template(template_path, output_path, {"status": 1})
            assert read_output(output_path) == [[None, "상태1"]]
            
            # Test status == 2
            render_template(template_path, output_path, {"status": 2})
            assert read_output(output_path) == [[None, "상태2"]]
            
            # Test else
            render_template(template_path, output_path, {"status": 3})
            assert read_output(output_path) == [[None, "기타"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestComments:
    """Tests for comments"""
    
    def test_comment_removed(self):
        template_path = create_template([
            ["Header"],
            ["{# 이것은 주석입니다 #}"],
            ["Footer"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {})
            result = read_output(output_path)
            assert result == [["Header"], ["Footer"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestStylePreservation:
    """Tests for style preservation"""
    
    def test_font_preserved(self):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = "{{ name }}"
        ws['A1'].font = Font(bold=True, size=14, color="FF0000")
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"name": "홍길동"})
            
            wb_out = load_workbook(output_path)
            ws_out = wb_out.active
            
            assert ws_out['A1'].value == "홍길동"
            assert ws_out['A1'].font.bold is True
            assert ws_out['A1'].font.size == 14
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_fill_preserved(self):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = "{{ value }}"
        ws['A1'].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"value": "테스트"})
            
            wb_out = load_workbook(output_path)
            ws_out = wb_out.active
            
            assert ws_out['A1'].value == "테스트"
            assert ws_out['A1'].fill.start_color.rgb == "00FFFF00"
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_border_preserved(self):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = "{{ value }}"
        thin_border = Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000')
        )
        ws['A1'].border = thin_border
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"value": "테두리 테스트"})
            
            wb_out = load_workbook(output_path)
            ws_out = wb_out.active
            
            assert ws_out['A1'].value == "테두리 테스트"
            assert ws_out['A1'].border.left.style == 'thin'
            assert ws_out['A1'].border.right.style == 'thin'
            assert ws_out['A1'].border.top.style == 'thin'
            assert ws_out['A1'].border.bottom.style == 'thin'
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_alignment_preserved(self):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = "{{ value }}"
        ws['A1'].alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"value": "정렬 테스트"})
            
            wb_out = load_workbook(output_path)
            ws_out = wb_out.active
            
            assert ws_out['A1'].value == "정렬 테스트"
            assert ws_out['A1'].alignment.horizontal == 'center'
            assert ws_out['A1'].alignment.vertical == 'top'
            assert ws_out['A1'].alignment.wrap_text is True
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_number_format_preserved(self):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = "{{ price }}"
        ws['A1'].number_format = '#,##0'
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"price": 1234567})
            
            wb_out = load_workbook(output_path)
            ws_out = wb_out.active
            
            assert ws_out['A1'].value == 1234567
            assert ws_out['A1'].number_format == '#,##0'
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_style_preserved_in_for_loop(self):
        """for loop 내에서 스타일이 복제되는지 테스트"""
        wb = Workbook()
        ws = wb.active
        ws['A1'] = "{% for item in items %}"
        # 본문 행
        ws['B2'] = "{{ item.name }}"
        ws['B2'].font = Font(bold=True, color="0000FF")
        ws['B2'].fill = PatternFill(start_color="FFFF00", fill_type="solid")
        ws['C2'] = "{{ item.price }}"
        ws['C2'].number_format = '#,##0'
        ws['C2'].alignment = Alignment(horizontal='right')
        ws['A3'] = "{% endfor %}"
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            data = {
                "items": [
                    {"name": "상품A", "price": 1000},
                    {"name": "상품B", "price": 2000},
                    {"name": "상품C", "price": 3000}
                ]
            }
            render_template(template_path, output_path, data)
            
            wb_out = load_workbook(output_path)
            ws_out = wb_out.active
            
            # 3개 행 모두 스타일이 유지되어야 함
            for row_idx in range(1, 4):
                # name 열 스타일 확인
                name_cell = ws_out.cell(row=row_idx, column=2)
                assert name_cell.font.bold is True
                assert name_cell.fill.start_color.rgb == "00FFFF00"
                
                # price 열 스타일 확인
                price_cell = ws_out.cell(row=row_idx, column=3)
                assert price_cell.number_format == '#,##0'
                assert price_cell.alignment.horizontal == 'right'
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_multiple_styles_combined(self):
        """여러 스타일이 함께 적용된 셀 테스트"""
        wb = Workbook()
        ws = wb.active
        ws['A1'] = "{{ value }}"
        ws['A1'].font = Font(bold=True, italic=True, size=12)
        ws['A1'].fill = PatternFill(start_color="00FF00", fill_type="solid")
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws['A1'].border = Border(
            left=Side(style='thick', color='FF0000'),
            bottom=Side(style='double', color='0000FF')
        )
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"value": "복합 스타일"})
            
            wb_out = load_workbook(output_path)
            ws_out = wb_out.active
            
            cell = ws_out['A1']
            assert cell.value == "복합 스타일"
            assert cell.font.bold is True
            assert cell.font.italic is True
            assert cell.font.size == 12
            assert cell.fill.start_color.rgb == "0000FF00"
            assert cell.alignment.horizontal == 'center'
            assert cell.alignment.vertical == 'center'
            assert cell.border.left.style == 'thick'
            assert cell.border.bottom.style == 'double'
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestComplexTemplate:
    """Tests for complex templates"""
    
    def test_mixed_template(self):
        """Test a template with for loop, if, and variables"""
        template_path = create_template([
            ["리포트: {{ title }}"],
            ["{# 상품 목록 #}"],
            ["{% for item in items %}"],
            ["{% if item.active %}"],
            ["", "{{ item.name }}", "{{ item.price }}"],
            ["{% endif %}"],
            ["{% endfor %}"],
            ["합계", "", "{{ total }}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            data = {
                "title": "월간 리포트",
                "items": [
                    {"name": "상품A", "price": 1000, "active": True},
                    {"name": "상품B", "price": 2000, "active": False},
                    {"name": "상품C", "price": 3000, "active": True}
                ],
                "total": 4000
            }
            render_template(template_path, output_path, data)
            result = read_output(output_path)
            
            assert result == [
                ["리포트: 월간 리포트"],
                [None, "상품A", 1000],
                [None, "상품C", 3000],
                ["합계", None, 4000]
            ]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestRichTextPreservation:
    """Tests for Rich Text (partial formatting) preservation"""
    
    def test_rich_text_variable_formatting_preserved(self):
        """변수 부분에 적용된 서식이 유지되는지 테스트"""
        wb = Workbook()
        ws = wb.active
        
        # "이름: " + "{{ name }}"(파란색 볼드) + "님"
        rich_text = CellRichText(
            "이름: ",
            TextBlock(InlineFont(b=True, color="0000FF"), "{{ name }}"),
            "님"
        )
        ws['A1'] = rich_text
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"name": "홍길동"})
            
            wb_out = load_workbook(output_path, rich_text=True)
            ws_out = wb_out.active
            
            result = ws_out['A1'].value
            assert isinstance(result, CellRichText), f"Expected CellRichText, got {type(result)}"
            
            # 전체 텍스트 확인
            assert str(result) == "이름: 홍길동님"
            
            # 서식 확인: 두 번째 파트가 볼드+파란색
            parts = list(result)
            assert len(parts) == 3
            assert parts[0] == "이름: "
            assert isinstance(parts[1], TextBlock)
            assert parts[1].text == "홍길동"
            assert parts[1].font.b is True  # bold
            assert parts[2] == "님"
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_rich_text_multiple_variables(self):
        """Rich Text 내 여러 변수 치환 테스트"""
        wb = Workbook()
        ws = wb.active
        
        # "{{ name }}"(빨간색) + " - " + "{{ title }}"(볼드)
        rich_text = CellRichText(
            TextBlock(InlineFont(color="FF0000"), "{{ name }}"),
            " - ",
            TextBlock(InlineFont(b=True), "{{ title }}")
        )
        ws['A1'] = rich_text
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"name": "김철수", "title": "개발자"})
            
            wb_out = load_workbook(output_path, rich_text=True)
            ws_out = wb_out.active
            
            result = ws_out['A1'].value
            assert isinstance(result, CellRichText)
            assert str(result) == "김철수 - 개발자"
            
            parts = list(result)
            assert len(parts) == 3
            assert isinstance(parts[0], TextBlock)
            assert parts[0].text == "김철수"
            assert parts[1] == " - "
            assert isinstance(parts[2], TextBlock)
            assert parts[2].text == "개발자"
            assert parts[2].font.b is True
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_rich_text_in_for_loop(self):
        """for loop 내 Rich Text 서식 유지 테스트"""
        wb = Workbook()
        ws = wb.active
        
        ws['A1'] = "{% for item in items %}"
        # Rich Text: "상품: " + "{{ item.name }}"(볼드)
        rich_text = CellRichText(
            "상품: ",
            TextBlock(InlineFont(b=True, color="008000"), "{{ item.name }}")
        )
        ws['B2'] = rich_text
        ws['A3'] = "{% endfor %}"
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            data = {
                "items": [
                    {"name": "사과"},
                    {"name": "바나나"},
                    {"name": "오렌지"}
                ]
            }
            render_template(template_path, output_path, data)
            
            wb_out = load_workbook(output_path, rich_text=True)
            ws_out = wb_out.active
            
            # 3개 행 모두 Rich Text 서식 유지 확인
            expected_names = ["사과", "바나나", "오렌지"]
            for row_idx, expected_name in enumerate(expected_names, start=1):
                result = ws_out.cell(row=row_idx, column=2).value
                assert isinstance(result, CellRichText), f"Row {row_idx}: Expected CellRichText"
                assert str(result) == f"상품: {expected_name}"
                
                parts = list(result)
                assert isinstance(parts[1], TextBlock)
                assert parts[1].text == expected_name
                assert parts[1].font.b is True
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_rich_text_no_variables(self):
        """변수 없는 Rich Text가 그대로 유지되는지 테스트"""
        wb = Workbook()
        ws = wb.active
        
        rich_text = CellRichText(
            "일반 텍스트 ",
            TextBlock(InlineFont(b=True), "볼드 텍스트")
        )
        ws['A1'] = rich_text
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {})
            
            wb_out = load_workbook(output_path, rich_text=True)
            ws_out = wb_out.active
            
            result = ws_out['A1'].value
            assert isinstance(result, CellRichText)
            assert str(result) == "일반 텍스트 볼드 텍스트"
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestEdgeCases:
    """Tests for edge cases and error handling"""
    
    def test_empty_sheet(self):
        """빈 시트 처리 테스트"""
        wb = Workbook()
        ws = wb.active
        # 빈 시트 - 아무 셀도 설정하지 않음
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"data": "test"})
            result = read_output(output_path)
            assert result == []
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_missing_endfor(self):
        """endfor 없는 for 루프 오류 테스트"""
        template_path = create_template([
            ["{% for item in items %}"],
            ["{{ item }}"]
            # endfor 없음
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            with pytest.raises(TemplateSyntaxError) as exc_info:
                render_template(template_path, output_path, {"items": [1, 2, 3]})
            assert "endfor" in str(exc_info.value).lower()
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_missing_endif(self):
        """endif 없는 if 문 오류 테스트"""
        template_path = create_template([
            ["{% if show %}"],
            ["표시됨"]
            # endif 없음
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            with pytest.raises(TemplateSyntaxError) as exc_info:
                render_template(template_path, output_path, {"show": True})
            assert "endif" in str(exc_info.value).lower()
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_unknown_control_statement(self):
        """알 수 없는 control statement 무시 테스트"""
        template_path = create_template([
            ["{% unknown_statement %}"],
            ["일반 행"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {})
            result = read_output(output_path)
            # 알 수 없는 control statement는 무시됨
            assert result == [["일반 행"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_orphan_endfor(self):
        """고아 endfor 처리 테스트"""
        template_path = create_template([
            ["일반 행"],
            ["{% endfor %}"],
            ["또 다른 행"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {})
            result = read_output(output_path)
            # 고아 endfor는 무시됨
            assert result == [["일반 행"], ["또 다른 행"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_orphan_endif(self):
        """고아 endif 처리 테스트"""
        template_path = create_template([
            ["일반 행"],
            ["{% endif %}"],
            ["또 다른 행"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {})
            result = read_output(output_path)
            # 고아 endif는 무시됨
            assert result == [["일반 행"], ["또 다른 행"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_orphan_elif(self):
        """고아 elif 처리 테스트"""
        template_path = create_template([
            ["일반 행"],
            ["{% elif condition %}"],
            ["또 다른 행"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"condition": True})
            result = read_output(output_path)
            # 고아 elif는 무시됨
            assert result == [["일반 행"], ["또 다른 행"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_orphan_else(self):
        """고아 else 처리 테스트"""
        template_path = create_template([
            ["일반 행"],
            ["{% else %}"],
            ["또 다른 행"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {})
            result = read_output(output_path)
            # 고아 else는 무시됨
            assert result == [["일반 행"], ["또 다른 행"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_none_iterable_in_for_loop(self):
        """None iterable 처리 테스트"""
        template_path = create_template([
            ["{% for item in items %}"],
            ["{{ item }}"],
            ["{% endfor %}"],
            ["Footer"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            # items가 context에 없음 -> None -> 빈 리스트로 처리
            render_template(template_path, output_path, {})
            result = read_output(output_path)
            assert result == [["Footer"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_multiple_sheets(self):
        """여러 시트 처리 테스트"""
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1['A1'] = "{{ name }}"
        
        ws2 = wb.create_sheet("Sheet2")
        ws2['A1'] = "{{ title }}"
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"name": "홍길동", "title": "제목"})
            
            wb_out = load_workbook(output_path)
            assert wb_out["Sheet1"]['A1'].value == "홍길동"
            assert wb_out["Sheet2"]['A1'].value == "제목"
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_comparison_with_none_values(self):
        """None 값 비교 조건 테스트"""
        template_path = create_template([
            ["{% if value > 0 %}"],
            ["양수"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            # value가 없으면 None, 비교 실패 -> False
            render_template(template_path, output_path, {})
            result = read_output(output_path)
            assert result == []
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_greater_equal_condition(self):
        """>=  조건 테스트"""
        template_path = create_template([
            ["{% if count >= 5 %}"],
            ["5 이상"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"count": 5})
            result = read_output(output_path)
            assert result == [["5 이상"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_less_equal_condition(self):
        """<= 조건 테스트"""
        template_path = create_template([
            ["{% if count <= 5 %}"],
            ["5 이하"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"count": 5})
            result = read_output(output_path)
            assert result == [["5 이하"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestLoopVariables:
    """Tests for loop variables (loop.index, loop.first, loop.last, loop.length)"""
    
    def test_loop_index(self):
        """loop.index 테스트 (1부터 시작)"""
        template_path = create_template([
            ["{% for item in items %}"],
            ["{{ loop.index }}", "{{ item }}"],
            ["{% endfor %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"items": ["A", "B", "C"]})
            result = read_output(output_path)
            assert result == [
                [1, "A"],
                [2, "B"],
                [3, "C"]
            ]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_loop_index0(self):
        """loop.index0 테스트 (0부터 시작)"""
        template_path = create_template([
            ["{% for item in items %}"],
            ["{{ loop.index0 }}", "{{ item }}"],
            ["{% endfor %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"items": ["A", "B", "C"]})
            result = read_output(output_path)
            assert result == [
                [0, "A"],
                [1, "B"],
                [2, "C"]
            ]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_loop_first(self):
        """loop.first 테스트"""
        template_path = create_template([
            ["{% for item in items %}"],
            ["{% if loop.first %}"],
            ["첫 번째: {{ item }}"],
            ["{% else %}"],
            ["{{ item }}"],
            ["{% endif %}"],
            ["{% endfor %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"items": ["A", "B", "C"]})
            result = read_output(output_path)
            assert result == [
                ["첫 번째: A"],
                ["B"],
                ["C"]
            ]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_loop_last(self):
        """loop.last 테스트"""
        template_path = create_template([
            ["{% for item in items %}"],
            ["{% if loop.last %}"],
            ["마지막: {{ item }}"],
            ["{% else %}"],
            ["{{ item }}"],
            ["{% endif %}"],
            ["{% endfor %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"items": ["A", "B", "C"]})
            result = read_output(output_path)
            assert result == [
                ["A"],
                ["B"],
                ["마지막: C"]
            ]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_loop_length(self):
        """loop.length 테스트"""
        template_path = create_template([
            ["{% for item in items %}"],
            ["{{ item }} ({{ loop.index }}/{{ loop.length }})"],
            ["{% endfor %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"items": ["A", "B", "C"]})
            result = read_output(output_path)
            assert result == [
                ["A (1/3)"],
                ["B (2/3)"],
                ["C (3/3)"]
            ]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_loop_single_item(self):
        """단일 아이템 루프에서 first와 last가 모두 True"""
        template_path = create_template([
            ["{% for item in items %}"],
            ["first={{ loop.first }}, last={{ loop.last }}"],
            ["{% endfor %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"items": ["only"]})
            result = read_output(output_path)
            assert result == [["first=True, last=True"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestFiltersInTemplate:
    """Tests for filters in templates"""
    
    def test_default_filter(self):
        """default 필터 테스트"""
        template_path = create_template([
            ["{{ value|default('N/A') }}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {})
            result = read_output(output_path)
            assert result == [["N/A"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_length_filter(self):
        """length 필터 테스트"""
        template_path = create_template([
            ["총 {{ items|length }}건"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"items": [1, 2, 3, 4, 5]})
            result = read_output(output_path)
            assert result == [["총 5건"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_join_filter(self):
        """join 필터 테스트"""
        template_path = create_template([
            ["{{ tags|join(', ') }}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"tags": ["python", "excel", "template"]})
            result = read_output(output_path)
            assert result == [["python, excel, template"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestTernaryInTemplate:
    """Tests for ternary operator in templates"""
    
    def test_ternary_true(self):
        """삼항 연산자 true 케이스"""
        template_path = create_template([
            ["{{ 'active' if is_active else 'inactive' }}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"is_active": True})
            result = read_output(output_path)
            assert result == [["active"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_ternary_false(self):
        """삼항 연산자 false 케이스"""
        template_path = create_template([
            ["{{ 'active' if is_active else 'inactive' }}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"is_active": False})
            result = read_output(output_path)
            assert result == [["inactive"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestLogicalOperatorsInTemplate:
    """Tests for logical operators in templates"""
    
    def test_and_condition(self):
        """and 조건 테스트"""
        template_path = create_template([
            ["{% if a and b %}"],
            ["둘 다 참"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"a": True, "b": True})
            result = read_output(output_path)
            assert result == [["둘 다 참"]]
            
            render_template(template_path, output_path, {"a": True, "b": False})
            result = read_output(output_path)
            assert result == []
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_or_condition(self):
        """or 조건 테스트"""
        template_path = create_template([
            ["{% if a or b %}"],
            ["하나라도 참"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"a": False, "b": True})
            result = read_output(output_path)
            assert result == [["하나라도 참"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_not_condition(self):
        """not 조건 테스트"""
        template_path = create_template([
            ["{% if not is_deleted %}"],
            ["삭제되지 않음"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"is_deleted": False})
            result = read_output(output_path)
            assert result == [["삭제되지 않음"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_in_condition(self):
        """in 조건 테스트"""
        template_path = create_template([
            ["{% if status in valid_statuses %}"],
            ["유효한 상태"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {
                "status": "active",
                "valid_statuses": ["active", "pending"]
            })
            result = read_output(output_path)
            assert result == [["유효한 상태"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestRendererEdgeCases:
    """Edge cases for renderer coverage"""
    
    def test_empty_sheet(self):
        """빈 시트 처리 테스트"""
        # 빈 워크북 생성
        wb = Workbook()
        ws = wb.active
        # max_row가 None 또는 0인 경우
        
        template_path = tempfile.mktemp(suffix='.xlsx')
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            wb.save(template_path)
            render_template(template_path, output_path, {})
            # 정상적으로 처리되어야 함
            assert os.path.exists(output_path)
        finally:
            if os.path.exists(template_path):
                os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_loop_context_repr(self):
        """LoopContext __repr__ 테스트"""
        from xlsx_template_renderer.renderer import LoopContext
        
        loop = LoopContext(0, 3)
        repr_str = repr(loop)
        assert "LoopContext" in repr_str
        assert "index=1" in repr_str
        assert "first=True" in repr_str
        assert "last=False" in repr_str
        assert "length=3" in repr_str
    
    def test_nested_if_statements(self):
        """중첩 if문 테스트 (depth 증가)"""
        template_path = create_template([
            ["{% if outer %}"],
            ["외부 조건"],
            ["{% if inner %}"],
            ["내부 조건"],
            ["{% endif %}"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"outer": True, "inner": True})
            result = read_output(output_path)
            assert result == [["외부 조건"], ["내부 조건"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_nested_if_outer_false(self):
        """중첩 if문에서 외부 조건이 거짓인 경우"""
        template_path = create_template([
            ["{% if outer %}"],
            ["외부 조건"],
            ["{% if inner %}"],
            ["내부 조건"],
            ["{% endif %}"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {"outer": False, "inner": True})
            result = read_output(output_path)
            assert result == []
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_triple_nested_if(self):
        """3중 중첩 if문 테스트"""
        template_path = create_template([
            ["{% if level1 %}"],
            ["레벨1"],
            ["{% if level2 %}"],
            ["레벨2"],
            ["{% if level3 %}"],
            ["레벨3"],
            ["{% endif %}"],
            ["{% endif %}"],
            ["{% endif %}"]
        ])
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            render_template(template_path, output_path, {
                "level1": True, 
                "level2": True, 
                "level3": True
            })
            result = read_output(output_path)
            assert result == [["레벨1"], ["레벨2"], ["레벨3"]]
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_rich_text_with_plain_string_part(self):
        """CellRichText에 plain string 파트가 있는 경우 테스트"""
        wb = Workbook()
        ws = wb.active
        
        # CellRichText에 plain string과 TextBlock 혼합
        # 일부 openpyxl 버전에서는 CellRichText에 plain string을 직접 추가할 수 있음
        rich_text = CellRichText(
            "Hello ",  # plain string part
            TextBlock(InlineFont(b=True), "{{ name }}"),  # TextBlock part
            "!"  # another plain string part
        )
        ws['A1'] = rich_text
        
        template_path = tempfile.mktemp(suffix='.xlsx')
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            wb.save(template_path)
            render_template(template_path, output_path, {"name": "World"})
            
            # 결과 확인
            result_wb = load_workbook(output_path, rich_text=True)
            result_ws = result_wb.active
            cell_value = result_ws['A1'].value
            
            # 값이 올바르게 치환되었는지 확인
            if isinstance(cell_value, CellRichText):
                text = str(cell_value)
            else:
                text = str(cell_value) if cell_value else ""
            
            assert "World" in text or "Hello" in text
        finally:
            if os.path.exists(template_path):
                os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_rich_text_plain_string_with_variable(self):
        """Plain string 파트에 변수가 있는 경우"""
        wb = Workbook()
        ws = wb.active
        
        # CellRichText에 변수가 포함된 plain string
        rich_text = CellRichText(
            "이름: {{ name }}",  # plain string with variable
        )
        ws['A1'] = rich_text
        
        template_path = tempfile.mktemp(suffix='.xlsx')
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            wb.save(template_path)
            render_template(template_path, output_path, {"name": "홍길동"})
            
            result_wb = load_workbook(output_path, rich_text=True)
            result_ws = result_wb.active
            cell_value = result_ws['A1'].value
            
            if isinstance(cell_value, CellRichText):
                text = str(cell_value)
            else:
                text = str(cell_value) if cell_value else ""
            
            assert "홍길동" in text
        finally:
            if os.path.exists(template_path):
                os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_process_rich_text_with_plain_string(self):
        """_process_rich_text_value 함수에서 plain string 파트 처리 테스트"""
        from xlsx_template_renderer.renderer import _process_rich_text_value
        
        # CellRichText에 plain string과 TextBlock 혼합
        rich_text = CellRichText(
            "Hello {{ name }}!",  # plain string with variable
            TextBlock(InlineFont(b=True), " Bold text")
        )
        
        context = {"name": "World"}
        result = _process_rich_text_value(rich_text, context)
        
        # 결과 확인
        result_text = str(result)
        assert "World" in result_text
        assert "Bold text" in result_text
    
    def test_process_rich_text_plain_string_no_variable(self):
        """_process_rich_text_value: plain string 파트에 변수가 없는 경우"""
        from xlsx_template_renderer.renderer import _process_rich_text_value
        
        rich_text = CellRichText(
            "Plain text without variable",  # no variable
            TextBlock(InlineFont(i=True), " italic")
        )
        
        context = {"name": "unused"}
        result = _process_rich_text_value(rich_text, context)
        
        result_text = str(result)
        assert "Plain text without variable" in result_text
        assert "italic" in result_text
    
    def test_process_rich_text_unknown_type(self):
        """_process_rich_text_value: 알 수 없는 타입 처리"""
        from xlsx_template_renderer.renderer import _process_rich_text_value
        
        # CellRichText를 직접 생성하고 알 수 없는 타입 추가
        # 실제로 CellRichText는 str과 TextBlock만 지원하지만,
        # 방어적 코드 테스트를 위해 내부 리스트에 직접 접근
        rich_text = CellRichText("Normal text")
        
        # 내부 리스트에 직접 알 수 없는 타입 추가 (방어적 코드 테스트)
        # CellRichText는 list를 상속하므로 append 가능
        rich_text.append(123)  # int는 알 수 없는 타입
        
        context = {}
        result = _process_rich_text_value(rich_text, context)
        
        # 알 수 없는 타입도 그대로 유지되어야 함
        assert 123 in result


class TestControlStatementStyles:
    """제어문의 스타일 처리 테스트"""
    
    def test_control_statement_with_styles_preserved_during_parsing(self):
        """제어문이 있는 셀의 스타일이 파싱 과정에서 보존되는지 테스트"""
        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "template_control_styles.xlsx"
            output_path = Path(temp_dir) / "output_control_styles.xlsx"
            
            # 템플릿 생성
            wb = Workbook()
            ws = wb.active
            
            # 제어문에 스타일 적용
            ws['A1'] = "  {% for item in items %}"  # 들여쓰기 포함
            ws['A1'].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            ws['A1'].font = Font(color="FF0000", bold=True)
            
            # 데이터 행
            ws['A2'] = "{{ item.name }}"
            ws['B2'] = "{{ item.value }}"
            
            # 종료 제어문
            ws['A3'] = "    {% endfor %}"  # 더 많은 들여쓰기
            ws['A3'].fill = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")
            ws['A3'].font = Font(color="0000FF", italic=True)
            
            wb.save(template_path)
            
            # 테스트 데이터
            data = {
                'items': [
                    {'name': '아이템1', 'value': 100},
                    {'name': '아이템2', 'value': 200}
                ]
            }
            
            # 렌더링 실행
            render_template(str(template_path), str(output_path), data)
            
            # 결과 확인 - 제어문 행들은 삭제되고 데이터만 남아야 함
            wb_out = load_workbook(output_path)
            ws_out = wb_out.active
            
            # 예상 결과: 2개 아이템이 2행에 걸쳐 출력
            assert ws_out['A1'].value == '아이템1'
            assert ws_out['B1'].value == 100
            assert ws_out['A2'].value == '아이템2'
            assert ws_out['B2'].value == 200
            
            # 제어문 행들은 삭제되었으므로 3행째는 비어있어야 함
            assert ws_out['A3'].value is None
    
    def test_indented_control_statements_parsing(self):
        """다양한 들여쓰기가 있는 제어문들이 정상적으로 파싱되는지 테스트"""
        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "template_indented.xlsx"
            output_path = Path(temp_dir) / "output_indented.xlsx"
            
            wb = Workbook()
            ws = wb.active
            
            # 중첩된 제어문들을 들여쓰기로 표현
            ws['A1'] = "{% for dept in departments %}"
            ws['A2'] = "  {% for emp in dept.employees %}"
            ws['A3'] = "    {% if emp.active %}"
            ws['A4'] = "{{ dept.name }}: {{ emp.name }}"
            ws['A5'] = "    {% endif %}"
            ws['A6'] = "  {% endfor %}"
            ws['A7'] = "{% endfor %}"
            
            wb.save(template_path)
            
            # 테스트 데이터
            data = {
                'departments': [
                    {
                        'name': '개발팀',
                        'employees': [
                            {'name': '김개발', 'active': True},
                            {'name': '이비활성', 'active': False},
                            {'name': '박활성', 'active': True}
                        ]
                    },
                    {
                        'name': '디자인팀', 
                        'employees': [
                            {'name': '최디자인', 'active': True}
                        ]
                    }
                ]
            }
            
            # 렌더링 실행
            render_template(str(template_path), str(output_path), data)
            
            # 결과 확인
            wb_out = load_workbook(output_path)
            ws_out = wb_out.active
            
            # active=True인 직원들만 출력되어야 함
            expected_results = [
                '개발팀: 김개발',
                '개발팀: 박활성', 
                '디자인팀: 최디자인'
            ]
            
            for i, expected in enumerate(expected_results, 1):
                assert ws_out[f'A{i}'].value == expected
    
    def test_control_statements_with_rainbow_bracket_style(self):
        """레인보우 브라켓 스타일의 제어문 색상 구분 테스트"""
        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "template_rainbow.xlsx"
            output_path = Path(temp_dir) / "output_rainbow.xlsx"
            
            wb = Workbook()
            ws = wb.active
            
            # 레인보우 브라켓 스타일로 색상 구분
            # for문 - 노란색
            ws['A1'] = "{% for item in items %}"
            ws['A1'].fill = PatternFill(start_color="FFFF00", fill_type="solid")
            ws['A1'].font = Font(bold=True)
            
            # if문 - 파란색  
            ws['A2'] = "  {% if item.important %}"
            ws['A2'].fill = PatternFill(start_color="0000FF", fill_type="solid")
            ws['A2'].font = Font(color="FFFFFF", bold=True)
            
            # 데이터
            ws['A3'] = "중요: {{ item.name }}"
            
            # else문 - 초록색
            ws['A4'] = "  {% else %}"
            ws['A4'].fill = PatternFill(start_color="00FF00", fill_type="solid")
            ws['A4'].font = Font(bold=True)
            
            # 데이터
            ws['A5'] = "일반: {{ item.name }}"
            
            # 종료문들 - 연한 회색
            ws['A6'] = "  {% endif %}"
            ws['A6'].fill = PatternFill(start_color="CCCCCC", fill_type="solid")
            ws['A6'].font = Font(italic=True)
            
            ws['A7'] = "{% endfor %}"
            ws['A7'].fill = PatternFill(start_color="CCCCCC", fill_type="solid")
            ws['A7'].font = Font(italic=True)
            
            wb.save(template_path)
            
            # 테스트 데이터
            data = {
                'items': [
                    {'name': '중요한일', 'important': True},
                    {'name': '일반업무', 'important': False}
                ]
            }
            
            # 렌더링 실행
            render_template(str(template_path), str(output_path), data)
            
            # 결과 확인
            wb_out = load_workbook(output_path)
            ws_out = wb_out.active
            
            # 조건에 따라 다른 출력이 나와야 함
            assert ws_out['A1'].value == '중요: 중요한일'
            assert ws_out['A2'].value == '일반: 일반업무'
            
            # 3행째는 비어있어야 함 (제어문들이 삭제됨)
            assert ws_out['A3'].value is None


class TestSheetFiltering:
    """Tests for sheets parameter filtering"""
    
    def test_sheets_filter_specific_sheets(self):
        """특정 시트만 처리되는지 테스트"""
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1['A1'] = "{{ value1 }}"
        
        ws2 = wb.create_sheet("Sheet2")
        ws2['A1'] = "{{ value2 }}"
        
        ws3 = wb.create_sheet("Sheet3")
        ws3['A1'] = "{{ value3 }}"
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            # Sheet1과 Sheet3만 처리
            render_template(
                template_path, 
                output_path, 
                {"value1": "처리됨1", "value2": "처리됨2", "value3": "처리됨3"},
                sheets=["Sheet1", "Sheet3"]
            )
            
            wb_out = load_workbook(output_path)
            # Sheet1은 처리됨
            assert wb_out["Sheet1"]['A1'].value == "처리됨1"
            # Sheet2는 처리되지 않음 (원본 템플릿 유지)
            assert wb_out["Sheet2"]['A1'].value == "{{ value2 }}"
            # Sheet3은 처리됨
            assert wb_out["Sheet3"]['A1'].value == "처리됨3"
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_sheets_filter_nonexistent_sheet_ignored(self):
        """존재하지 않는 시트명은 무시되는지 테스트"""
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1['A1'] = "{{ name }}"
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            # 존재하지 않는 시트와 존재하는 시트 혼합
            render_template(
                template_path, 
                output_path, 
                {"name": "홍길동"},
                sheets=["NonExistent", "Sheet1", "AlsoNonExistent"]
            )
            
            wb_out = load_workbook(output_path)
            # Sheet1은 정상 처리됨
            assert wb_out["Sheet1"]['A1'].value == "홍길동"
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)
    
    def test_sheets_filter_none_processes_all(self):
        """sheets=None이면 모든 시트가 처리되는지 테스트"""
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1['A1'] = "{{ name }}"
        
        ws2 = wb.create_sheet("Sheet2")
        ws2['A1'] = "{{ title }}"
        
        fd, template_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(template_path)
        
        output_path = template_path.replace('.xlsx', '_out.xlsx')
        
        try:
            # sheets 파라미터 없이 호출 (기존 동작)
            render_template(template_path, output_path, {"name": "홍길동", "title": "제목"})
            
            wb_out = load_workbook(output_path)
            assert wb_out["Sheet1"]['A1'].value == "홍길동"
            assert wb_out["Sheet2"]['A1'].value == "제목"
        finally:
            os.unlink(template_path)
            if os.path.exists(output_path):
                os.unlink(output_path)

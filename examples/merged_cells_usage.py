#!/usr/bin/env python3
"""
병합 셀 템플릿 사용 예제

이 예제는 for 루프 내에서 병합 셀이 어떻게 복제되는지 보여줍니다.
- 수직 병합 (여러 행)
- 수평 병합 (여러 열)
- 병합 셀 내 여러 변수
"""

from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from xlsx_template_renderer import render_template_to_file

EXAMPLES_DIR = Path(__file__).parent


def create_sample_template():
    """병합 셀이 포함된 샘플 템플릿 생성"""
    wb = Workbook()
    ws = wb.active
    ws.title = "주문목록"
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 헤더 (수평 병합 + 여러 변수)
    ws['A1'] = '{{ report_title }} - {{ report_date }}'
    ws.merge_cells('A1:D1')
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')
    
    # 열 헤더
    ws['A2'] = '카테고리/코드'
    ws['B2'] = '상품명'
    ws['C2'] = '수량'
    ws['D2'] = '가격'
    for col in ['A', 'B', 'C', 'D']:
        ws[f'{col}2'].border = thin_border
        ws[f'{col}2'].fill = PatternFill(start_color='DDDDDD', fill_type='solid')
    
    # for 루프 시작
    ws['A3'] = '{% for item in items %}'
    
    # 데이터 행 1: 카테고리 (수직 병합 A4:A5 + 여러 변수)
    ws['A4'] = '{{ item.category }} ({{ item.category_code }})'
    ws['B4'] = '{{ item.name }}'
    ws['C4'] = '{{ item.quantity }}'
    ws['D4'] = '{{ item.unit_price }}'
    
    # 데이터 행 2: 소계 (수평 병합 B5:C5 + 여러 변수)
    ws['A5'] = ''  # 수직 병합의 일부
    ws['B5'] = '{{ item.name }} 소계: {{ item.quantity }}개'
    ws['D5'] = '{{ item.subtotal }}'
    
    # for 루프 종료
    ws['A6'] = '{% endfor %}'
    
    # 수직 병합 (카테고리)
    ws.merge_cells('A4:A5')
    ws['A4'].alignment = Alignment(vertical='center', horizontal='center')
    
    # 수평 병합 (소계 레이블)
    ws.merge_cells('B5:C5')
    ws['B5'].alignment = Alignment(horizontal='right')
    
    # 스타일 적용
    for row in range(4, 6):
        for col in ['A', 'B', 'C', 'D']:
            ws[f'{col}{row}'].border = thin_border
    
    # 푸터 (수평 병합 + 여러 변수)
    ws['A7'] = '총 {{ items|length }}건, 합계:'
    ws.merge_cells('A7:C7')
    ws['A7'].alignment = Alignment(horizontal='right')
    ws['A7'].font = Font(bold=True)
    ws['D7'] = '{{ total }}'
    ws['D7'].font = Font(bold=True)
    ws['D7'].border = thin_border
    
    # 열 너비 조정
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    
    # 저장
    template_path = EXAMPLES_DIR / "merged_cells_template.xlsx"
    wb.save(template_path)
    print(f"템플릿 생성됨: {template_path}")
    return template_path


def main():
    # 템플릿 생성
    template_path = create_sample_template()
    output_path = EXAMPLES_DIR / "merged_cells_output.xlsx"
    
    # 샘플 데이터
    data = {
        "report_title": "2024년 1분기 주문 현황",
        "report_date": "2024-03-31",
        "items": [
            {
                "category": "전자제품",
                "category_code": "EL",
                "name": "노트북",
                "quantity": 2,
                "unit_price": 1500000,
                "subtotal": 3000000,
            },
            {
                "category": "가전",
                "category_code": "HA",
                "name": "냉장고",
                "quantity": 1,
                "unit_price": 800000,
                "subtotal": 800000,
            },
            {
                "category": "가구",
                "category_code": "FN",
                "name": "책상",
                "quantity": 3,
                "unit_price": 200000,
                "subtotal": 600000,
            },
        ],
        "total": 4400000,
    }
    
    # 렌더링 실행
    render_template_to_file(str(template_path), str(output_path), data)
    print(f"렌더링 완료: {output_path}")
    
    # 결과 확인
    wb = load_workbook(output_path)
    ws = wb.active
    
    print("\n=== 렌더링 결과 ===")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        print(row)
    
    print(f"\n병합된 셀 범위: {[str(r) for r in ws.merged_cells.ranges]}")


if __name__ == "__main__":
    main()

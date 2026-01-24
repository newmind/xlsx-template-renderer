"""
예제: include_section을 사용한 재사용 가능한 템플릿 컴포넌트

이 예제에서는:
- define_section으로 재사용 가능한 섹션 정의
- include_section으로 정의된 섹션 삽입
- for 루프 내에서 include_section 사용
- if 조건문과 include_section 조합
"""

from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

from xlsx_template_renderer import render_template


def create_include_section_template():
    """include_section을 사용하는 멀티 시트 템플릿 생성"""
    wb = Workbook()
    
    # ===== 메인 시트 =====
    ws_main = wb.active
    ws_main.title = "메인"
    
    # 스타일 정의
    title_font = Font(bold=True, size=14, color="1F4E79")
    
    # 헤더 섹션 포함
    ws_main['A1'] = '{% include_section "컴포넌트" "헤더" %}'
    
    # 본문 제목
    ws_main['A2'] = "주문 목록"
    ws_main['A2'].font = title_font
    
    # 주문별 반복
    ws_main['A3'] = '{% for order in orders %}'
    
    # 조건에 따라 다른 섹션 사용
    ws_main['A4'] = '{% if order.is_vip %}'
    ws_main['A5'] = '  {% include_section "컴포넌트" "VIP_주문" %}'
    ws_main['A6'] = '{% else %}'
    ws_main['A7'] = '  {% include_section "컴포넌트" "일반_주문" %}'
    ws_main['A8'] = '{% endif %}'
    
    ws_main['A9'] = '{% endfor %}'
    
    # 푸터 섹션 포함
    ws_main['A10'] = '{% include_section "컴포넌트" "푸터" %}'
    
    # ===== 컴포넌트 시트 =====
    ws_comp = wb.create_sheet("컴포넌트")
    
    # 헤더 섹션 정의
    ws_comp['A1'] = '{# define_section:헤더 #}'
    ws_comp['A2'] = "{{ company_name }}"
    ws_comp['A2'].font = Font(bold=True, size=16)
    ws_comp['A3'] = "작성일: {{ report_date }}"
    ws_comp['A4'] = ""
    ws_comp['A5'] = '{# enddefine_section #}'
    
    # VIP 주문 섹션 정의 (섹션 내에서 for 사용)
    ws_comp['A7'] = '{# define_section:VIP_주문 #}'
    ws_comp['A8'] = "⭐ VIP 주문 #{{ order.id }}"
    ws_comp['A8'].font = Font(bold=True, color="FFD700")
    ws_comp['A8'].fill = PatternFill(start_color="1F4E79", fill_type="solid")
    ws_comp['B8'] = "{{ order.customer }}"
    ws_comp['C8'] = "{{ order.amount }}"
    ws_comp['C8'].number_format = '#,##0원'
    # 섹션 내부에서 for 루프 - 주문 상품 목록
    ws_comp['A9'] = '{% for item in order.items %}'
    ws_comp['A10'] = "  - {{ item.name }}"
    ws_comp['B10'] = "{{ item.qty }}개"
    ws_comp['C10'] = "{{ item.price }}"
    ws_comp['C10'].number_format = '#,##0원'
    ws_comp['A11'] = '{% endfor %}'
    ws_comp['A12'] = '{# enddefine_section #}'
    
    # 일반 주문 섹션 정의 (섹션 내에서 for 사용)
    ws_comp['A14'] = '{# define_section:일반_주문 #}'
    ws_comp['A15'] = "주문 #{{ order.id }}"
    ws_comp['B15'] = "{{ order.customer }}"
    ws_comp['C15'] = "{{ order.amount }}"
    ws_comp['C15'].number_format = '#,##0원'
    # 섹션 내부에서 for 루프 - 주문 상품 목록
    ws_comp['A16'] = '{% for item in order.items %}'
    ws_comp['A17'] = "  - {{ item.name }}"
    ws_comp['B17'] = "{{ item.qty }}개"
    ws_comp['C17'] = "{{ item.price }}"
    ws_comp['C17'].number_format = '#,##0원'
    ws_comp['A18'] = '{% endfor %}'
    ws_comp['A19'] = '{# enddefine_section #}'
    
    # 푸터 섹션 정의
    ws_comp['A21'] = '{# define_section:푸터 #}'
    ws_comp['A22'] = ""
    ws_comp['A23'] = "총 {{ orders|length }}건의 주문"
    ws_comp['A23'].font = Font(italic=True, color="666666")
    ws_comp['A24'] = "문의: {{ contact_email }}"
    ws_comp['A24'].font = Font(color="0066CC")
    ws_comp['A25'] = '{# enddefine_section #}'
    
    # 열 너비 조정
    ws_main.column_dimensions['A'].width = 25
    ws_main.column_dimensions['B'].width = 15
    ws_main.column_dimensions['C'].width = 15
    
    ws_comp.column_dimensions['A'].width = 25
    ws_comp.column_dimensions['B'].width = 15
    ws_comp.column_dimensions['C'].width = 15
    
    # 저장
    template_path = Path(__file__).parent / "include_section_template.xlsx"
    wb.save(template_path)
    print(f"템플릿 생성됨: {template_path}")
    return template_path


def main():
    # 템플릿 생성
    template_path = create_include_section_template()
    output_path = Path(__file__).parent / "include_section_output.xlsx"
    
    # 렌더링할 데이터 (주문별 items로 중첩 데이터 테스트)
    data = {
        "company_name": "메이크스타 쇼핑몰",
        "report_date": "2024-01-20",
        "contact_email": "support@makestar.com",
        "orders": [
            {"id": 1001, "customer": "김철수", "amount": 150000, "is_vip": True,
             "items": [
                 {"name": "프리미엄 앨범", "qty": 2, "price": 50000},
                 {"name": "포토카드 세트", "qty": 1, "price": 50000},
             ]},
            {"id": 1002, "customer": "이영희", "amount": 35000, "is_vip": False,
             "items": [
                 {"name": "일반 앨범", "qty": 1, "price": 35000},
             ]},
            {"id": 1003, "customer": "박지민", "amount": 280000, "is_vip": True,
             "items": [
                 {"name": "한정판 앨범", "qty": 1, "price": 150000},
                 {"name": "굿즈 패키지", "qty": 2, "price": 65000},
             ]},
            {"id": 1004, "customer": "최수진", "amount": 42000, "is_vip": False,
             "items": [
                 {"name": "포스터 세트", "qty": 3, "price": 14000},
             ]},
        ]
    }
    
    # 렌더링 실행 (메인 시트만 처리, 컴포넌트 시트는 참조용)
    render_template(str(template_path), str(output_path), data, sheets=["메인"])
    print(f"렌더링 완료: {output_path}")
    
    # 결과 확인
    wb = load_workbook(output_path)
    ws = wb["메인"]
    
    print("\n" + "=" * 50)
    print("렌더링 결과")
    print("=" * 50)
    
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        row_str = [str(cell) if cell else "" for cell in row]
        if any(row_str):
            print(f"  {row_str}")
    
    print("\n" + "=" * 50)
    print("테스트된 기능:")
    print("  - define_section/enddefine_section으로 섹션 정의")
    print("  - include_section으로 다른 시트의 섹션 삽입")
    print("  - for 루프 내에서 include_section 사용")
    print("  - if 조건문과 include_section 조합")
    print("  - 섹션 내부에서 for 루프 사용 (order.items 순회)")
    print("  - 중첩 데이터 context 전달 (order -> order.items)")
    print("  - 스타일이 포함된 섹션 복사")
    print("  - 변수 치환 (섹션 내부의 {{ }} 처리)")


if __name__ == "__main__":
    main()

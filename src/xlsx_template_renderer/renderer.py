"""
Main renderer for xlsx templates.
Processes template files and outputs rendered xlsx files with styles preserved.
"""

import copy
from typing import Any, Dict, List, Optional, Tuple
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from .parser import (
    parse_cell, 
    is_control_statement, 
    extract_variables, 
    substitute_variables,
    evaluate_condition,
    Token, 
    TokenType
)
from .expressions import evaluate_expression, resolve_path
from .exceptions import TemplateSyntaxError, TemplateRenderError


class LoopContext:
    """
    Provides loop metadata for {% for %} loops.
    
    Attributes:
        index: 1-based iteration count
        index0: 0-based iteration count
        first: True if first iteration
        last: True if last iteration
        length: Total number of items
    """
    
    def __init__(self, index0: int, length: int):
        self.index = index0 + 1
        self.index0 = index0
        self.first = index0 == 0
        self.last = index0 == length - 1
        self.length = length
    
    def __repr__(self):
        return f"LoopContext(index={self.index}, first={self.first}, last={self.last}, length={self.length})"


def render_template(
    template_path: str,
    output_path: str,
    data: Dict[str, Any]
) -> None:
    """
    Render an xlsx template with the given data.
    
    Args:
        template_path: Path to the template xlsx file
        output_path: Path to save the rendered xlsx file
        data: Dictionary containing the data to render
    """
    # Load the template (rich_text=True to preserve Rich Text formatting)
    wb = load_workbook(template_path, rich_text=True)
    
    # Process each sheet
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        _render_sheet(ws, data)
    
    # Save the result
    wb.save(output_path)


def _render_sheet(ws: Worksheet, data: Dict[str, Any]) -> None:
    """
    Render a single worksheet.
    
    Strategy:
    1. Parse all rows and identify control structures
    2. Build a list of output rows with context
    3. Write output rows to a new structure
    4. Delete/shift rows as needed
    """
    # Collect all row data first
    max_row = ws.max_row
    max_col = ws.max_column
    
    if max_row is None or max_row == 0:
        return
    
    # Parse the template structure
    rows_data = []
    for row_idx in range(1, max_row + 1):
        row_cells = []
        for col_idx in range(1, (max_col or 0) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            row_cells.append({
                'value': cell.value,
                'style': _copy_cell_style(cell),
            })
        rows_data.append(row_cells)
    
    # Process template and generate output rows
    output_rows = _process_rows(rows_data, data)
    
    # Clear the worksheet
    for row_idx in range(1, max_row + 1):
        for col_idx in range(1, (max_col or 0) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = None
    
    # Write output rows
    for out_row_idx, row_data in enumerate(output_rows, start=1):
        for col_idx, cell_data in enumerate(row_data, start=1):
            cell = ws.cell(row=out_row_idx, column=col_idx)
            cell.value = cell_data['value']
            _apply_cell_style(cell, cell_data['style'])
    
    # Delete extra rows if output is shorter
    if len(output_rows) < max_row:
        ws.delete_rows(len(output_rows) + 1, max_row - len(output_rows))


def _copy_cell_style(cell: Cell) -> dict:
    """Copy cell style properties."""
    return {
        'font': copy.copy(cell.font),
        'fill': copy.copy(cell.fill),
        'border': copy.copy(cell.border),
        'alignment': copy.copy(cell.alignment),
        'number_format': cell.number_format,
        'protection': copy.copy(cell.protection),
    }


def _apply_cell_style(cell: Cell, style: dict) -> None:
    """Apply style properties to a cell."""
    if style.get('font'):
        cell.font = style['font']
    if style.get('fill'):
        cell.fill = style['fill']
    if style.get('border'):
        cell.border = style['border']
    if style.get('alignment'):
        cell.alignment = style['alignment']
    if style.get('number_format'):
        cell.number_format = style['number_format']
    if style.get('protection'):
        cell.protection = style['protection']


def _process_rows(
    rows_data: List[List[dict]], 
    context: Dict[str, Any]
) -> List[List[dict]]:
    """
    Process template rows and return output rows.
    
    Handles for loops, if statements, and variable substitution.
    """
    output_rows = []
    row_idx = 0
    
    while row_idx < len(rows_data):
        row = rows_data[row_idx]
        first_cell_value = row[0]['value'] if row else None
        
        # Check if this is a control statement row
        if first_cell_value and is_control_statement(str(first_cell_value)):
            token = parse_cell(str(first_cell_value))
            
            if token is None:
                # Unknown control statement, skip the row
                row_idx += 1
                continue
            
            if token.type == TokenType.COMMENT:
                # Skip comment rows
                row_idx += 1
                continue
            
            elif token.type == TokenType.FOR_START:
                # Process for loop
                end_idx = _find_matching_end(rows_data, row_idx, TokenType.FOR_START, TokenType.FOR_END)
                if end_idx is None:
                    raise TemplateSyntaxError(f"Missing {{% endfor %}} for loop starting", row=row_idx + 1)
                
                # Get the iterable
                iterable = resolve_path(context, token.loop_iter)
                if iterable is None:
                    iterable = []
                
                # Convert to list to get length
                iterable_list = list(iterable) if not isinstance(iterable, (list, tuple)) else iterable
                iterable_length = len(iterable_list)
                
                # Process loop body for each item
                body_rows = rows_data[row_idx + 1:end_idx]
                
                for idx, item in enumerate(iterable_list):
                    # Create loop context object
                    loop_obj = LoopContext(idx, iterable_length)
                    # Create new context with loop variable and loop object
                    loop_context = {**context, token.loop_var: item, 'loop': loop_obj}
                    # Recursively process body rows
                    processed = _process_rows(body_rows, loop_context)
                    output_rows.extend(processed)
                
                # Skip to after endfor
                row_idx = end_idx + 1
                continue
            
            elif token.type == TokenType.IF_START:
                # Process if statement
                end_idx, branches = _find_if_branches(rows_data, row_idx)
                if end_idx is None:
                    raise TemplateSyntaxError(f"Missing {{% endif %}} for if starting", row=row_idx + 1)
                
                # Evaluate branches in order
                executed = False
                for branch_type, condition, start, end in branches:
                    if branch_type == 'if' or branch_type == 'elif':
                        if evaluate_condition(condition, context):
                            body_rows = rows_data[start:end]
                            processed = _process_rows(body_rows, context)
                            output_rows.extend(processed)
                            executed = True
                            break
                    elif branch_type == 'else':
                        if not executed:
                            body_rows = rows_data[start:end]
                            processed = _process_rows(body_rows, context)
                            output_rows.extend(processed)
                        break
                
                # Skip to after endif
                row_idx = end_idx + 1
                continue
            
            elif token.type in (TokenType.FOR_END, TokenType.IF_END, TokenType.ELIF, TokenType.ELSE):
                # These should be handled by their parent constructs
                row_idx += 1
                continue
        
        # Regular row - process variables and add to output
        processed_row = _process_row_variables(row, context)
        output_rows.append(processed_row)
        row_idx += 1
    
    return output_rows


def _process_row_variables(row: List[dict], context: Dict[str, Any]) -> List[dict]:
    """Process variable substitutions in a row."""
    result = []
    
    for cell_data in row:
        value = cell_data['value']
        new_value = value
        
        if value is not None:
            # Handle Rich Text (CellRichText)
            if isinstance(value, CellRichText):
                new_value = _process_rich_text_value(value, context)
            else:
                str_value = str(value)
                variables = extract_variables(str_value)
                
                if variables:
                    # Build replacements dict
                    replacements = {}
                    for expr in variables:
                        evaluated = evaluate_expression(expr, context)
                        if evaluated is not None:
                            replacements[expr] = evaluated
                    
                    # Substitute variables
                    new_value = substitute_variables(str_value, replacements)
                    
                    # If the entire cell was a single variable, preserve the type
                    if str_value.strip().startswith('{{') and str_value.strip().endswith('}}'):
                        single_expr = variables[0] if len(variables) == 1 else None
                        if single_expr and single_expr in replacements:
                            # Use the actual value type (number, etc.)
                            new_value = replacements[single_expr]
        
        result.append({
            'value': new_value,
            'style': cell_data['style'],
        })
    
    return result


def _process_rich_text_value(rich_text: CellRichText, context: Dict[str, Any]) -> CellRichText:
    """
    Process variable substitutions in Rich Text while preserving formatting.
    
    Args:
        rich_text: CellRichText object containing text and formatting
        context: Data context for variable evaluation
    
    Returns:
        New CellRichText with variables substituted
    """
    result_parts = []
    
    for part in rich_text:
        if isinstance(part, TextBlock):
            # TextBlock has font and text
            text = part.text
            font = copy.copy(part.font)
            
            variables = extract_variables(text)
            if variables:
                # Build replacements
                replacements = {}
                for expr in variables:
                    evaluated = evaluate_expression(expr, context)
                    if evaluated is not None:
                        replacements[expr] = str(evaluated)
                
                # Substitute variables
                new_text = substitute_variables(text, replacements)
                result_parts.append(TextBlock(font, new_text))
            else:
                result_parts.append(TextBlock(font, text))
        elif isinstance(part, str):
            # Plain string part
            variables = extract_variables(part)
            if variables:
                replacements = {}
                for expr in variables:
                    evaluated = evaluate_expression(expr, context)
                    if evaluated is not None:
                        replacements[expr] = str(evaluated)
                
                new_text = substitute_variables(part, replacements)
                result_parts.append(new_text)
            else:
                result_parts.append(part)
        else:
            # Unknown type, keep as is
            result_parts.append(part)
    
    return CellRichText(result_parts)


def _find_matching_end(
    rows_data: List[List[dict]], 
    start_idx: int, 
    start_type: TokenType, 
    end_type: TokenType
) -> Optional[int]:
    """Find the matching end tag for a control structure."""
    depth = 0
    
    for idx in range(start_idx, len(rows_data)):
        row = rows_data[idx]
        first_cell_value = row[0]['value'] if row else None
        
        if first_cell_value and is_control_statement(str(first_cell_value)):
            token = parse_cell(str(first_cell_value))
            if token:
                if token.type == start_type:
                    depth += 1
                elif token.type == end_type:
                    depth -= 1
                    if depth == 0:
                        return idx
    
    return None


def _find_if_branches(
    rows_data: List[List[dict]], 
    start_idx: int
) -> Tuple[Optional[int], List[Tuple[str, str, int, int]]]:
    """
    Find all branches of an if statement.
    
    Returns:
        (endif_idx, [(branch_type, condition, body_start, body_end), ...])
    """
    branches = []
    depth = 0
    current_branch_start = start_idx + 1
    current_branch_type = 'if'
    current_condition = ''
    
    # Get the initial if condition
    first_row = rows_data[start_idx]
    if first_row:
        token = parse_cell(str(first_row[0]['value']))
        if token and token.type == TokenType.IF_START:
            current_condition = token.expression
    
    for idx in range(start_idx, len(rows_data)):
        row = rows_data[idx]
        first_cell_value = row[0]['value'] if row else None
        
        if first_cell_value and is_control_statement(str(first_cell_value)):
            token = parse_cell(str(first_cell_value))
            if token:
                if token.type == TokenType.IF_START:
                    if idx == start_idx:
                        depth = 1
                    else:
                        depth += 1
                
                elif token.type == TokenType.ELIF and depth == 1:
                    # End current branch, start elif branch
                    branches.append((current_branch_type, current_condition, current_branch_start, idx))
                    current_branch_type = 'elif'
                    current_condition = token.expression
                    current_branch_start = idx + 1
                
                elif token.type == TokenType.ELSE and depth == 1:
                    # End current branch, start else branch
                    branches.append((current_branch_type, current_condition, current_branch_start, idx))
                    current_branch_type = 'else'
                    current_condition = ''
                    current_branch_start = idx + 1
                
                elif token.type == TokenType.IF_END:
                    depth -= 1
                    if depth == 0:
                        # End the last branch
                        branches.append((current_branch_type, current_condition, current_branch_start, idx))
                        return idx, branches
    
    return None, branches
